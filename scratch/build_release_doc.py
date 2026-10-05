from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()

    # Título principal
    title = doc.add_heading('Documento de Liberación (Release) - Helarate', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Changelog
    doc.add_heading('1. Documentación de Release (Changelog)', level=1)
    doc.add_paragraph('Qué se libera: Primera versión estable (v1.0.0) de la aplicación web Helarate. Se incluye el sistema principal de ventas y gestión.', style='List Bullet')
    doc.add_paragraph('Qué cambió: Implementación de interfaz de usuario con Flutter Web, integración de base de datos Supabase, y sistema de autenticación.', style='List Bullet')
    doc.add_paragraph('Aprobado por: Equipo de Desarrollo e Ingeniería.', style='List Bullet')

    # 2. Selección de herramienta
    doc.add_heading('2. Selección justificada de herramienta de despliegue', level=1)
    p = doc.add_paragraph()
    p.add_run('Herramienta seleccionada: ').bold = True
    p.add_run('Vercel.\n')
    p.add_run('Justificación: ').bold = True
    p.add_run('Se evaluaron alternativas como GitHub Pages y Firebase Hosting. Vercel fue elegido por su integración continua (CI/CD) nativa y automática con el repositorio de GitHub, su facilidad para desplegar aplicaciones web modernas y tiempos de compilación muy rápidos. Además, requiere nula configuración de servidores, lo cual es ideal para un proyecto frontend como Flutter Web.')

    # 3. Publicación accesible
    doc.add_heading('3. Publicación accesible (URL Funcional)', level=1)
    doc.add_paragraph('La aplicación se encuentra desplegada y es accesible en la siguiente URL:')
    # Placeholder URL, user can update
    p_url = doc.add_paragraph('https://helarate.vercel.app', style='Intense Quote')
    p_url.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 4. Evidencia del proceso de despliegue
    doc.add_heading('4. Evidencia del proceso de despliegue', level=1)
    doc.add_paragraph('A continuación, se adjunta la evidencia (capturas del dashboard de Vercel y logs de compilación) que demuestra el proceso automatizado de publicación:')
    doc.add_paragraph('[INSERTAR CAPTURAS DE PANTALLA DE VERCEL/GITHUB ACTIONS AQUÍ]')
    # Add a blank box/space for screenshots
    doc.add_paragraph('\n\n\n')

    # 5. Plan de Rollback
    doc.add_heading('5. Plan de Rollback', level=1)
    p_roll = doc.add_paragraph()
    p_roll.add_run('Procedimiento de reversión en caso de fallos críticos en producción:\n').bold = True
    doc.add_paragraph('1. En caso de detectarse un fallo crítico, el equipo ingresará al panel de control de Vercel.', style='List Number')
    doc.add_paragraph('2. En la sección "Deployments" (Despliegues), se buscará el despliegue anterior que funcionaba correctamente (última versión estable).', style='List Number')
    doc.add_paragraph('3. Se seleccionará la opción "Promote to Production" o "Instant Rollback" (Revertir al instante) para que el tráfico web sea redirigido inmediatamente a la versión anterior sin necesidad de compilar nuevamente.', style='List Number')
    doc.add_paragraph('4. Se analizará y corregirá el código problemático en una rama de desarrollo separada, probándose antes de generar un nuevo release.', style='List Number')

    file_name = 'DOCUMENTO_LIBERACION_HELARATE.docx'
    doc.save(file_name)
    print(f'Archivo creado con éxito: {file_name}')

if __name__ == '__main__':
    main()
